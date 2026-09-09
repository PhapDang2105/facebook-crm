from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill


PROJECT_ROOT = Path(__file__).resolve().parents[2]
SOURCE = PROJECT_ROOT / "references" / "quan-ly-doanh-nghiep" / "backend" / "assets" / "export-template.xlsx"
OUTPUT = PROJECT_ROOT / "assets" / "templates" / "facebook-order-export.xlsx"
DATA_COLUMN_COUNT = 43


workbook = load_workbook(SOURCE)
worksheet = workbook["FileNhapDonHang"]

for row in range(1, 5):
    for column in range(1, DATA_COLUMN_COUNT + 1):
        cell = worksheet.cell(row, column)
        font = copy(cell.font)
        font.name = "Times New Roman"
        cell.font = font

data_styles = [copy(worksheet.cell(4, column)._style) for column in range(1, DATA_COLUMN_COUNT + 1)]
data_height = worksheet.row_dimensions[4].height

for sheet in list(workbook.worksheets):
    if sheet is not worksheet:
        workbook.remove(sheet)

if worksheet.max_row >= 4:
    worksheet.delete_rows(4, worksheet.max_row - 3)

worksheet.data_validations.dataValidation = []
worksheet.auto_filter.ref = None
worksheet.sheet_view.selection[0].activeCell = "A1"
worksheet.sheet_view.selection[0].sqref = "A1"

for column, style in enumerate(data_styles, start=1):
    cell = worksheet.cell(4, column)
    cell._style = style
    cell.fill = PatternFill(fill_type=None)
    cell.value = None

worksheet.row_dimensions[4].height = data_height
workbook.active = 0
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
workbook.save(OUTPUT)
print(OUTPUT)
